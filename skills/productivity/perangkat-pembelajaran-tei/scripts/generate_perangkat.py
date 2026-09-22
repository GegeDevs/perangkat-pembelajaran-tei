#!/root/.hermes/skills/cv-ats-format/.venv/bin/python
from pathlib import Path
import sys
import struct, zlib
# lib_cv memuat Pillow/pypdf untuk fitur foto dan hitung PDF; keduanya tidak
# diperlukan pada perangkat pembelajaran teks ini. Sediakan shim import agar
# tetap memakai fungsi layout dan konversi resmi dari skill cv-ats-format.
try:
    import PIL  # noqa: F401
except ImportError:
    import types
    pil = types.ModuleType('PIL'); pil.Image = types.ModuleType('PIL.Image'); pil.ImageOps = types.ModuleType('PIL.ImageOps')
    sys.modules['PIL'] = pil; sys.modules['PIL.Image'] = pil.Image; sys.modules['PIL.ImageOps'] = pil.ImageOps
try:
    import pypdf  # noqa: F401
except ImportError:
    import types
    pp = types.ModuleType('pypdf'); pp.PdfReader = object; sys.modules['pypdf'] = pp
sys.path.insert(0, '/root/.hermes/skills/cv-ats-format/scripts')
from lib_cv import set_doc_style, add_section_title, add_text, add_bullet, add_rich_paragraph, convert_docx_to_pdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/root/hermes/supervisi-rc-25sept')
OUT.mkdir(parents=True, exist_ok=True)
FONT='Times New Roman'
TEACHER='Gege Desembri, S.Pd.'
NIP='199712302025211018'
IMG = OUT / 'ilustrasi_resistor_gelang_warna.png'
CIRCUIT_IMG = OUT / 'ilustrasi_rangkaian_seri_paralel.png'
CAPACITOR_IMG = OUT / 'ilustrasi_kapasitor_polar_nonpolar.png'
CAPACITOR_EXAMPLE_IMG = OUT / 'ilustrasi_contoh_gelang_resistor.png'

COLOR_HEX = {
    'Hitam': '000000', 'Cokelat': '653219', 'Merah': 'C82323', 'Jingga': 'F28C28',
    'Kuning': 'FFD92F', 'Hijau': '2E8B57', 'Biru': '3B82D0', 'Ungu': '7B3FB5',
    'Abu-abu': 'A6A6A6', 'Putih': 'FFFFFF', 'Emas': 'D9B52E', 'Perak': 'C0C0C0'
}

RESISTOR_EXAMPLES = [
    ('Cokelat–Hitam–Merah–Emas', '10 × 10² Ω ±5%', '1 kΩ ±5%', [(101,50,25),(0,0,0),(190,35,35),(218,180,45)]),
    ('Merah–Merah–Cokelat–Emas', '22 × 10¹ Ω ±5%', '220 Ω ±5%', [(190,35,35),(190,35,35),(101,50,25),(218,180,45)]),
    ('Kuning–Ungu–Jingga–Emas', '47 × 10³ Ω ±5%', '47 kΩ ±5%', [(240,210,35),(123,63,181),(242,140,40),(218,180,45)])
]

# Ilustrasi gabungan rangkaian resistor dan kapasitor seri-paralel.
def make_circuit_png(path):
    import subprocess, textwrap
    script = textwrap.dedent(r'''
        from PIL import Image, ImageDraw, ImageFont
        import sys
        out, font_path = sys.argv[1], sys.argv[2]
        im=Image.new('RGB',(1400,900),'white'); d=ImageDraw.Draw(im)
        font=ImageFont.truetype(font_path,28); small=ImageFont.truetype(font_path,24); bold=ImageFont.truetype(font_path,34)
        def wire(points): d.line(points,fill=(35,35,35),width=6,joint='curve')
        def resistor(x,y,label):
            d.rectangle((x,y-28,x+150,y+28),fill=(232,183,105),outline=(35,35,35),width=4)
            d.text((x+42,y-18),label,font=small,fill=(20,20,20))
        def capacitor(x,y,label):
            # Kabel masuk menyentuh pelat pertama; kabel keluar dimulai dari pelat kedua.
            d.line((x,y,x+65,y),fill=(35,35,35),width=6)
            d.line((x+65,y-55,x+65,y+55),fill=(35,35,35),width=7)
            d.line((x+95,y-55,x+95,y+55),fill=(35,35,35),width=7)
            d.text((x+8,y+78),label,font=small,fill=(20,20,20))
        def node(x,y): d.ellipse((x-9,y-9,x+9,y+9),fill=(35,35,35))
        # headings and panels
        d.text((45,25),'Rangkaian Resistor',font=bold,fill=(0,0,0)); d.text((745,25),'Rangkaian Kapasitor',font=bold,fill=(0,0,0))
        d.text((45,90),'SERI',font=bold,fill=(30,80,150)); d.text((745,90),'SERI',font=bold,fill=(30,80,150))
        y=170; wire([(55,y),(190,y)]); resistor(190,y,'R1'); wire([(340,y),(420,y)]); resistor(420,y,'R2'); wire([(570,y),(680,y)]); node(55,y); node(680,y)
        wire([(755,y),(870,y)]); capacitor(870,y,'C1'); wire([(965,y),(1045,y)]); capacitor(1045,y,'C2'); wire([(1140,y),(1340,y)]); node(755,y); node(1340,y)
        d.text((45,270),'Rs = R1 + R2',font=font,fill=(0,0,0)); d.text((745,270),'1/Cs = 1/C1 + 1/C2',font=font,fill=(0,0,0))
        d.text((45,390),'PARALEL',font=bold,fill=(30,120,70)); d.text((745,390),'PARALEL',font=bold,fill=(30,120,70))
        # resistor parallel branches
        x0,x1=55,680; yt,yb=490,650; wire([(x0,yt),(x0,yb)]); wire([(x1,yt),(x1,yb)]); node(x0,570); node(x1,570)
        wire([(x0,yt),(190,yt)]); resistor(190,yt,'R1'); wire([(340,yt),(x1,yt)])
        wire([(x0,yb),(190,yb)]); resistor(190,yb,'R2'); wire([(340,yb),(x1,yb)])
        # capacitor parallel branches
        x0,x1=755,1340; wire([(x0,yt),(x0,yb)]); wire([(x1,yt),(x1,yb)]); node(x0,570); node(x1,570)
        wire([(x0,yt),(870,yt)]); capacitor(870,yt,'C1'); wire([(965,yt),(x1,yt)])
        wire([(x0,yb),(870,yb)]); capacitor(870,yb,'C2'); wire([(965,yb),(x1,yb)])
        d.text((45,790),'1/Rp = 1/R1 + 1/R2',font=font,fill=(0,0,0)); d.text((745,790),'Cp = C1 + C2',font=font,fill=(0,0,0))
        im.save(out,'PNG')
    ''')
    font = subprocess.check_output(['fc-match','-f','%{file}','Arial'], text=True).strip()
    subprocess.run(['/usr/bin/python3','-c',script,str(path),font],check=True)


def _png_chunk(kind, data):
    return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)

# Font bitmap bold sans-serif 5x7: stroke tebal dan bentuk huruf sederhana.
# Dirancang khusus agar tetap terbaca pada gelang yang sempit.
FONT5 = {
    'A':['01110','11011','11011','11111','11011','11011','11011'], 'B':['11110','11011','11011','11110','11011','11011','11110'],
    'C':['01111','11000','11000','11000','11000','11000','01111'], 'D':['11110','11011','11011','11011','11011','11011','11110'],
    'E':['11111','11000','11000','11110','11000','11000','11111'], 'G':['01111','11000','11000','11011','11011','11011','01111'],
    'H':['11011','11011','11011','11111','11011','11011','11011'], 'I':['11111','01100','01100','01100','01100','01100','11111'],
    'K':['11011','11011','11110','11100','11110','11011','11011'], 'L':['11000','11000','11000','11000','11000','11000','11111'],
    'M':['11011','11111','11111','11011','11011','11011','11011'], 'N':['11011','11111','11111','11111','11011','11011','11011'],
    'O':['01110','11011','11011','11011','11011','11011','01110'], 'R':['11110','11011','11011','11110','11100','11011','11011'],
    'T':['11111','01100','01100','01100','01100','01100','01100'], 'U':['11011','11011','11011','11011','11011','11011','01110'],
}

def draw_text(pix, text, x, y, scale=4, color=(20,20,20), spacing=2):
    # Font bitmap kecil agar teks dapat digambar tanpa Pillow.
    width=sum((5 if ch!=' ' else 3)*scale+spacing for ch in text)-spacing
    for ch in text:
        if ch==' ':
            x += 3*scale+spacing; continue
        glyph=FONT5.get(ch, ['00000']*7)
        for gy,row in enumerate(glyph):
            for gx,val in enumerate(row):
                if val=='1':
                    for yy in range(y+gy*scale,y+(gy+1)*scale):
                        for xx in range(x+gx*scale,x+(gx+1)*scale):
                            if 0<=yy<len(pix) and 0<=xx<len(pix[0]): pix[yy][xx]=list(color)
        x += 5*scale+spacing
    return width

def draw_text_vertical(pix, text, x, y, scale=3, color=(255,255,255), spacing=1):
    # Menulis label vertikal agar seluruh nama warna berada di dalam gelang.
    glyphs=[FONT5.get(ch, ['00000']*7) for ch in text]
    char_h=7*scale
    for i,glyph in enumerate(glyphs):
        oy=y+i*(char_h+spacing)
        for gy,row in enumerate(glyph):
            for gx,val in enumerate(row):
                if val=='1':
                    for yy in range(oy+gy*scale,oy+(gy+1)*scale):
                        for xx in range(x+gx*scale,x+(gx+1)*scale):
                            if 0<=yy<len(pix) and 0<=xx<len(pix[0]): pix[yy][xx]=list(color)

def make_resistor_png(path):
    # Ilustrasi raster PNG tanpa dependensi eksternal: resistor 4 gelang.
    w, h = 1200, 430
    bg=(255,255,255); pix=[[list(bg) for _ in range(w)] for _ in range(h)]
    def rect(x0,y0,x1,y1,c):
        for y in range(max(0,y0),min(h,y1)):
            for x in range(max(0,x0),min(w,x1)): pix[y][x]=list(c)
    def line(x0,y0,x1,y1,c,th=6):
        if x0==x1: rect(x0-th//2,y0,x1+th//2,y1,c)
        else: rect(x0,y0-th//2,x1,y1+th//2,c)
    # terminal wires and end caps
    line(80,215,290,215,(80,80,80),12); line(910,215,1120,215,(80,80,80),12)
    rect(270,120,930,310,(232,183,105)); rect(270,120,930,140,(150,105,50)); rect(270,290,930,310,(150,105,50))
    # Ujung badan dibuat polos tanpa garis tepi agar tidak disalahartikan sebagai gelang.
    colors=[(101,50,25),(0,0,0),(190,35,35),(218,180,45)]
    for x,c in zip([390,500,610,790],colors): rect(x,135,x+54,295,c)
    # small black outline-like markers around band positions
    for x in [390,500,610,790]:
        rect(x,135,x+4,295,(35,35,35)); rect(x+50,135,x+54,295,(35,35,35))
    # Label warna ditempatkan langsung di dalam gelang dan ditulis vertikal.
    # Label emas memakai tinta hitam agar kontras; gelang lain memakai putih.
    labels=[('COKELAT',390,(255,255,255)),('HITAM',500,(255,255,255)),('MERAH',610,(255,255,255)),('EMAS',790,(20,20,20))]
    for label,x,ink in labels:
        tw=5*3
        draw_text_vertical(pix,label,x+27-tw//2,141,scale=3,color=ink,spacing=1)
    # encode RGB PNG
    raw=b''.join(b'\x00'+bytes(sum(row,[])) for row in pix)
    png=b'\x89PNG\r\n\x1a\n'+_png_chunk(b'IHDR',struct.pack('>IIBBBBB',w,h,8,2,0,0,0))+_png_chunk(b'IDAT',zlib.compress(raw,9))+_png_chunk(b'IEND',b'')
    path.write_bytes(png)

def make_resistor_png(path):
    # Render label menggunakan Arial jika tersedia; fallback Linux yang kompatibel
    # adalah Liberation Sans (metriks dan bentuk sans-serif serupa Arial).
    import subprocess
    import textwrap
    script = textwrap.dedent(r'''
        from PIL import Image, ImageDraw, ImageFont
        import sys
        out=sys.argv[1]
        W,H=1200,430
        im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
        d.line((80,215,290,215),fill=(80,80,80),width=12); d.line((910,215,1120,215),fill=(80,80,80),width=12)
        d.rectangle((270,120,930,310),fill=(232,183,105))
        d.rectangle((270,120,930,140),fill=(150,105,50)); d.rectangle((270,290,930,310),fill=(150,105,50))
        colors=[(101,50,25),(0,0,0),(190,35,35),(218,180,45)]
        for x,c in zip((390,500,610,790),colors):
            d.rectangle((x,135,x+54,295),fill=c)
        # Garis tipis hanya membatasi gelang, bukan ujung body resistor.
        for x in (390,500,610,790):
            d.rectangle((x,135,x+4,295),fill=(35,35,35)); d.rectangle((x+50,135,x+54,295),fill=(35,35,35))
        font_path=sys.argv[2]
        font=ImageFont.truetype(font_path,22)
        labels=[('COKELAT',390,(255,255,255)),('HITAM',500,(255,255,255)),('MERAH',610,(255,255,255)),('EMAS',790,(20,20,20))]
        for label,x,ink in labels:
            # Teks diputar 90 derajat dan diletakkan tepat di dalam gelang.
            layer=Image.new('RGBA',(90,250),(0,0,0,0)); ld=ImageDraw.Draw(layer)
            box=ld.textbbox((0,0),label,font=font,stroke_width=0); tw=box[2]-box[0]; th=box[3]-box[1]
            ld.text(((90-tw)//2,(250-th)//2),label,font=font,fill=ink)
            layer=layer.rotate(90,expand=True)
            im.paste(layer,(x+27-layer.width//2,215-layer.height//2),layer)
        im.save(out,'PNG')
    ''')
    font = subprocess.check_output(['fc-match','-f','%{file}','Arial'], text=True).strip()
    subprocess.run(['/usr/bin/python3','-c',script,str(path),font],check=True)

def add_image(doc, path, width_cm=14, caption=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); r.add_picture(str(path), width=Cm(width_cm))
    caption = caption or 'Ilustrasi: resistor 4 gelang warna — cokelat, hitam, merah, emas.'
    cap=doc.add_paragraph(caption); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs: r.font.name=FONT; r.font.size=Pt(10); r.font.italic=True


def add_circuit_image(doc):
    add_image(doc, CIRCUIT_IMG, width_cm=16,
              caption='Ilustrasi: rangkaian resistor dan kapasitor seri-paralel beserta rumus nilai penggantinya.')

def make_example_resistors(path):
    import subprocess, textwrap
    script = textwrap.dedent(r'''
        from PIL import Image, ImageDraw, ImageFont
        import sys
        out, font_path = sys.argv[1], sys.argv[2]
        im=Image.new('RGB',(1050,300),'white'); d=ImageDraw.Draw(im)
        font=ImageFont.truetype(font_path,22)
        examples=[((101,50,25),(0,0,0),(190,35,35),(218,180,45),'1 kΩ ±5%'),((190,35,35),(190,35,35),(101,50,25),(218,180,45),'220 Ω ±5%'),((240,210,35),(123,63,181),(242,140,40),(218,180,45),'47 kΩ ±5%')]
        for i,(a,b,c,e,label) in enumerate(examples):
            x=30+i*345; y=95
            d.line((x,y+35,x+35,y+35),fill=(70,70,70),width=6); d.line((x+245,y+35,x+315,y+35),fill=(70,70,70),width=6)
            d.rounded_rectangle((x+35,y,x+245,y+70),radius=12,fill=(232,183,105),outline=(50,50,50),width=3)
            for bx,col in zip((x+75,x+115,x+155,x+205),(a,b,c,e)): d.rectangle((bx,y+3,bx+18,y+67),fill=col)
            d.text((x+47,y+85),label,font=font,fill=(0,0,0))
        im.save(out,'PNG')
    ''')
    font=subprocess.check_output(['fc-match','-f','%{file}','Liberation Sans'],text=True).strip()
    subprocess.run(['/usr/bin/python3','-c',script,str(path),font],check=True)

def make_capacitor_png(path):
    import subprocess, textwrap
    script=textwrap.dedent(r'''
        from PIL import Image, ImageDraw, ImageFont
        import sys
        out,font_path=sys.argv[1],sys.argv[2]
        im=Image.new('RGB',(1200,520),'white'); d=ImageDraw.Draw(im)
        f=ImageFont.truetype(font_path,28); s=ImageFont.truetype(font_path,24)
        def text(x,y,t,ft=f): d.text((x,y),t,font=ft,fill=(0,0,0))
        def symbol(x,y,polar):
            d.line((x,y+90,x+75,y+90),fill=(35,35,35),width=6); d.line((x+75,y+35,x+75,y+145),fill=(35,35,35),width=8)
            d.line((x+105,y+35,x+105,y+145),fill=(35,35,35),width=8); d.line((x+105,y+90,x+180,y+90),fill=(35,35,35),width=6)
            if polar: text(x+128,y+25,'+',s)
        text(90,35,'Kapasitor Polar (Elektrolit)'); text(700,35,'Kapasitor Non-Polar')
        # component bodies
        d.rounded_rectangle((95,230,280,390),radius=25,fill=(45,105,170),outline=(25,25,25),width=4); text(135,275,'+   C',f); text(112,410,'Ada tanda + dan −',s)
        d.rounded_rectangle((705,230,890,390),radius=25,fill=(210,175,75),outline=(25,25,25),width=4); text(755,290,'C',f); text(712,410,'Tidak berpolaritas',s)
        text(90,140,'Simbol rangkaian:',s); symbol(270,120,True); text(700,140,'Simbol rangkaian:',s); symbol(880,120,False)
        im.save(out,'PNG')
    ''')
    font=subprocess.check_output(['fc-match','-f','%{file}','Liberation Sans'],text=True).strip()
    subprocess.run(['/usr/bin/python3','-c',script,str(path),font],check=True)

def add_capacitor_image(doc):
    add_image(doc,CAPACITOR_IMG,width_cm=16,caption='Ilustrasi: komponen kapasitor polar dan non-polar beserta simbolnya.')


def make_support_images():
    make_example_resistors(CAPACITOR_EXAMPLE_IMG)
    make_capacitor_png(CAPACITOR_IMG)

def style_doc(doc):
    set_doc_style(doc, 21.0, 29.7, margin_cm=2.0, body_size=12)
    for s in doc.styles:
        if s.type == 1:
            s.font.name=FONT
            s._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for sec in doc.sections:
        sec.header_distance=Cm(0.8); sec.footer_distance=Cm(0.8)
        hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=hp.add_run('SMK NEGERI 1 TERUSAN NUNYAI | DASAR-DASAR TEKNIK ELEKTRONIKA')
        r.font.name=FONT; r.font.size=Pt(9); r.font.bold=True
        fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=fp.add_run('Perangkat Pembelajaran — Gege Desembri, S.Pd. — 25 September 2026')
        r.font.name=FONT; r.font.size=Pt(9)

def title(doc, text, subtitle=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.font.name=FONT; r.font.size=Pt(14); r.font.bold=True
    if subtitle:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(subtitle); r.font.name=FONT; r.font.size=Pt(12); r.font.bold=True

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill); shd.set(qn('w:val'),'clear')

def table(doc, headers, rows, widths=None, color_column=None, color_map=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs: r.font.name=FONT; r.font.size=Pt(11); r.font.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if color_column is not None and i == color_column and color_map:
                fill=color_map.get(str(val))
                if fill: set_cell_shading(cells[i],fill)
                if str(val) in ('Hitam','Cokelat','Merah','Biru','Ungu'):
                    for p in cells[i].paragraphs:
                        for r in p.runs: r.font.color.rgb=__import__('docx').shared.RGBColor(255,255,255)
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.name=FONT; r.font.size=Pt(11)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    return t

def section(doc, text): add_section_title(doc, text, space_before=8, space_after=3, font_size=12)
def para(doc,text,**kw): add_text(doc,text,**kw)
def bullets(doc, items):
    for x in items: add_bullet(doc,x,font_size=11)
def numbered(doc, items):
    for i,x in enumerate(items,1): para(doc,f'{i}. {x}',size=11)
def save(doc,name):
    path=OUT/name; doc.save(path)
    pdf=convert_docx_to_pdf(path, OUT)
    return path,pdf

def common_identity(doc, title_text):
    title(doc,title_text,'KURIKULUM MERDEKA — FASE E')
    table(doc,['Komponen','Keterangan'],[
      ['Satuan Pendidikan','SMK Negeri 1 Terusan Nunyai'],['Program/Konsentrasi','Teknik Elektronika / Teknik Elektronika Industri'],
      ['Mata Pelajaran','Dasar-Dasar Teknik Elektronika'],['Fase/Kelas','E / X (Sepuluh)'],['Semester','1 (Ganjil)'],
      ['Materi','Resistor dan Kapasitor'],['Alokasi Waktu','1 pertemuan × 6 JP × 40 menit = 240 menit'],['Guru','Gege Desembri, S.Pd. — NIP '+NIP],['Tanggal','Jumat, 25 September 2026']
    ],[4.5,11.5])

def make_rpp():
    d=Document(); style_doc(d); common_identity(d,'RENCANA PELAKSANAAN PEMBELAJARAN (RPP)')
    section(d,'A. CAPAIAN PEMBELAJARAN')
    para(d,'Pada akhir Fase E, peserta didik mampu memahami komponen elektronika pasif dan aktif, termasuk jenis, karakteristik, fungsi, simbol, bentuk/kemasan, serta membaca sistem kode komponen sesuai standar dan penerapannya dalam rangkaian elektronika.')
    section(d,'B. TUJUAN PEMBELAJARAN')
    numbered(d,['Mengidentifikasi simbol, fungsi, jenis, dan karakteristik resistor serta kapasitor.','Membaca kode warna resistor dan menentukan nilai resistansi beserta toleransinya.','Membaca kode angka/huruf kapasitor serta mengonversi satuan pF, nF, dan µF.','Menghitung nilai total resistor dan kapasitor pada rangkaian seri dan paralel.','Menyelesaikan masalah sederhana komponen R dan C melalui kerja kelompok dengan teliti, mandiri, dan bertanggung jawab.'])
    section(d,'C. KETERKAITAN ATP DAN PROFIL PELAJAR PANCASILA')
    para(d,'Elemen 9: Komponen Elektronika Aktif dan Pasif. Acuan TP: DTE 9.1 (jenis, bentuk/kemasan, dan karakteristik komponen pasif) serta DTE 9.3 (pembacaan kode nilai/sistem kode komponen pasif).')
    para(d,'Profil yang dikembangkan: bernalar kritis, mandiri, gotong royong, dan kreatif.')
    section(d,'D. MODEL, PENDEKATAN, METODE, MEDIA, DAN SUMBER')
    table(d,['Aspek','Rincian'],[['Model','Problem Based Learning (PBL)'],['Pendekatan','Saintifik dan kontekstual'],['Metode','Demonstrasi, diskusi, praktik membaca komponen, latihan, presentasi, refleksi'],['Media/Alat','Komponen resistor dan kapasitor, LKPD, papan tulis/LCD, kalkulator, multimeter bila tersedia'],['Sumber','Bahan ajar guru, buku Dasar-Dasar Teknik Elektronika SMK Kelas X, datasheet/label komponen']],[4.5,11.5])
    section(d,'E. LANGKAH-LANGKAH PEMBELAJARAN (240 MENIT)')
    table(d,['Tahap','Aktivitas Pembelajaran','Waktu'],[
      ['Pendahuluan','Salam, doa, presensi, apersepsi dengan menunjukkan resistor dan kapasitor, penyampaian tujuan, manfaat, aturan kerja, dan asesmen.','20 menit'],
      ['Fase 1 — Orientasi masalah','Guru menyajikan masalah: bagaimana menentukan nilai komponen kecil dan memilih susunan R/C yang tepat ketika label terbatas? Peserta didik mengamati dan mengajukan dugaan.','25 menit'],
      ['Fase 2 — Organisasi','Peserta didik dibagi menjadi kelompok, menerima LKPD dan sampel komponen, membagi peran pencatat, pembaca kode, penghitung, dan penyaji.','15 menit'],
      ['Fase 3 — Penyelidikan','Kelompok mengidentifikasi simbol/fungsi, membaca kode warna resistor, kode kapasitor, konversi satuan, serta menghitung rangkaian seri-paralel R dan C. Guru membimbing.','95 menit'],
      ['Fase 4 — Penyajian','Kelompok menyajikan satu hasil pembacaan kode dan satu hasil perhitungan. Kelompok lain memberi tanggapan berdasarkan LKPD.','35 menit'],
      ['Fase 5 — Evaluasi','Guru dan peserta didik memeriksa jawaban, menegaskan perbedaan rumus resistor dan kapasitor, serta membahas kesalahan umum.','20 menit'],
      ['Penutup','Kesimpulan, refleksi 3-2-1, tes individu singkat, umpan balik, tindak lanjut, doa, dan salam.','30 menit']
    ],[3.2,10.8,2])
    section(d,'F. ASESMEN')
    table(d,['Aspek','Teknik/Instrumen','Kriteria/Ketuntasan'],[['Sikap','Observasi: tanggung jawab, kerja sama, ketelitian, kemandirian','Minimal kategori Baik'],['Pengetahuan','Tes individu: kode R/C dan perhitungan seri-paralel','Nilai minimal 75'],['Keterampilan','Unjuk kerja membaca komponen dan produk LKPD','Minimal skor 3 pada sebagian besar indikator']],[3,8,5])
    section(d,'G. PENGAYAAN, REMEDIAL, DAN REFLEKSI')
    bullets(d,['Pengayaan: menganalisis rangkaian campuran R-C dan membaca kode kapasitor dengan toleransi/tegangan kerja berbeda.','Remedial: menggunakan kartu warna dan tabel konversi satuan, contoh bertahap, serta latihan terbimbing sampai mencapai nilai minimal.','Refleksi guru: mencatat miskonsepsi tentang kapasitor seri-paralel dan strategi pembimbingan kelompok.'])
    section(d,'H. PENGESAHAN')
    para(d,'Mengetahui,\nKepala SMK Negeri 1 Terusan Nunyai\n\n\n( ________________________ )\nNIP. ______________________',align=WD_ALIGN_PARAGRAPH.LEFT)
    para(d,'Terusan Nunyai, 25 September 2026\nGuru Mata Pelajaran,\n\n\nGege Desembri, S.Pd.\nNIP. '+NIP,align=WD_ALIGN_PARAGRAPH.RIGHT)
    return save(d,'RPP_ResistorKapasitor_X-TEI_GegeDesembri.docx')

def make_bahan():
    d=Document(); style_doc(d); title(d,'BAHAN AJAR','PERHITUNGAN RESISTOR DAN KAPASITOR — KELAS X TEI')
    para(d,'Guru: '+TEACHER+' | SMKN 1 Terusan Nunyai | Alokasi: 6 JP × 40 menit')
    section(d,'PENDAHULUAN — TUJUAN PEMBELAJARAN')
    bullets(d,['Mengenali fungsi, simbol, jenis, dan karakteristik resistor serta kapasitor.','Membaca nilai resistor dari kode warna dan kapasitor dari kode angka/huruf.','Menghitung nilai total resistor dan kapasitor seri-paralel.'])
    section(d,'BAB I — TEORI DASAR KELISTRIKAN')
    section(d,'2.1 Hukum Ohm')
    para(d,'Hukum Ohm menyatakan bahwa arus listrik yang mengalir pada suatu penghantar berbanding lurus dengan tegangan dan berbanding terbalik dengan hambatan, selama kondisi fisik penghantar tetap. Rumusnya V = I × R, sehingga I = V/R dan R = V/I. V menyatakan tegangan dalam volt (V), I menyatakan arus dalam ampere (A), dan R menyatakan hambatan dalam ohm (Ω).')
    table(d,['Besaran','Simbol','Satuan','Rumus'],[['Tegangan','V','volt (V)','V = I × R'],['Arus','I','ampere (A)','I = V/R'],['Hambatan','R','ohm (Ω)','R = V/I']],[4,3,4,5])
    para(d,'Contoh: sumber 12 V dipasang pada resistor 1 kΩ. Arus yang mengalir adalah I = 12/1.000 = 0,012 A = 12 mA.')
    section(d,'2.2 Daya Listrik pada Resistor')
    para(d,'Daya listrik adalah laju penggunaan energi listrik. Pada resistor, daya dapat dihitung dengan P = V × I. Dengan menggabungkan Hukum Ohm, rumusnya dapat ditulis P = I² × R atau P = V²/R. Satuan daya adalah watt (W). Pilih resistor dengan daya nominal lebih besar daripada daya yang diperkirakan agar komponen tidak panas berlebihan.')
    para(d,'Contoh: resistor 1 kΩ dialiri arus 12 mA. P = I² × R = (0,012)² × 1.000 = 0,144 W. Resistor ¼ W masih memenuhi secara teoritis, tetapi tetap perlu mempertimbangkan faktor keamanan dan suhu kerja.')
    section(d,'2.3 Hukum Kirchhoff secara Pengantar')
    para(d,'Hukum Arus Kirchhoff (KCL) menyatakan bahwa jumlah arus yang masuk ke suatu titik percabangan sama dengan jumlah arus yang keluar. Hukum Tegangan Kirchhoff (KVL) menyatakan bahwa jumlah aljabar tegangan dalam satu loop tertutup sama dengan nol. Kedua hukum ini membantu menganalisis rangkaian seri, paralel, dan rangkaian yang lebih kompleks.')
    section(d,'2.4 Kapasitansi, Muatan, dan Energi')
    para(d,'Kapasitansi menunjukkan kemampuan kapasitor menyimpan muatan. Hubungannya dinyatakan dengan Q = C × V, dengan Q dalam coulomb (C), C dalam farad (F), dan V dalam volt (V). Energi yang tersimpan pada kapasitor adalah E = ½ × C × V² dengan satuan joule (J). Jangan menyentuh kapasitor yang masih bermuatan; lakukan pelepasan muatan melalui cara yang aman.')
    section(d,'2.5 Pengisian Kapasitor dan Konstanta Waktu RC')
    para(d,'Pada rangkaian resistor-kapasitor, konstanta waktu ditentukan oleh τ = R × C. Setelah satu konstanta waktu, tegangan kapasitor saat pengisian mencapai sekitar 63% dari tegangan sumber; setelah sekitar lima konstanta waktu, kapasitor dianggap hampir penuh. Saat pelepasan, tegangan turun hingga sekitar 37% setelah satu konstanta waktu. Nilai R dalam ohm dan C dalam farad.')
    section(d,'BAB II — RESISTOR')
    add_image(d, IMG)
    section(d,'2.1 Pengertian dan Fungsi Resistor')
    para(d,'Resistor adalah komponen pasif yang memberikan hambatan terhadap arus listrik. Satuan resistansi adalah ohm (Ω). Fungsi resistor antara lain membatasi arus, membagi tegangan, dan menentukan kondisi kerja rangkaian.')
    section(d,'2.2 Kode Warna dan Tabel Gelang Resistor')
    para(d,'Kode warna 4 gelang: gelang 1 dan 2 = angka signifikan; gelang 3 = pengali; gelang 4 = toleransi. Baca gelang dari sisi yang paling dekat dengan ujung resistor; gelang toleransi biasanya berjarak lebih lebar.')
    table(d,['Warna','Angka','Pengali','Toleransi umum'],[['Hitam','0','×10⁰','—'],['Cokelat','1','×10¹','±1%'],['Merah','2','×10²','±2%'],['Jingga','3','×10³','—'],['Kuning','4','×10⁴','—'],['Hijau','5','×10⁵','±0,5%'],['Biru','6','×10⁶','±0,25%'],['Ungu','7','×10⁷','±0,1%'],['Abu-abu','8','×10⁸','—'],['Putih','9','×10⁹','—'],['Emas','—','×10⁻¹','±5%'],['Perak','—','×10⁻²','±10%']],[4,3,4,5],color_column=0,color_map=COLOR_HEX)
    para(d,'Cara membaca: tentukan dua angka pertama, kalikan dengan faktor gelang ketiga, lalu tuliskan toleransi gelang keempat. Contoh cokelat–hitam–merah–emas = 10 × 10² Ω ±5% = 1 kΩ ±5%.')
    add_image(d, CAPACITOR_EXAMPLE_IMG, width_cm=16, caption='Ilustrasi contoh gelang resistor dan nilai resistansinya.')
    table(d,['Contoh gelang','Perhitungan','Nilai'],[['Cokelat–Hitam–Merah–Emas','10 × 10² Ω ±5%','1 kΩ ±5%'],['Merah–Merah–Cokelat–Emas','22 × 10¹ Ω ±5%','220 Ω ±5%'],['Kuning–Ungu–Jingga–Emas','47 × 10³ Ω ±5%','47 kΩ ±5%']],[5,7,4])
    section(d,'BAB III — KAPASITOR')
    section(d,'3.1 Pengertian, Fungsi, dan Satuan')
    para(d,'Kapasitor adalah komponen pasif yang tersusun atas dua konduktor (pelat elektroda) yang dipisahkan oleh bahan isolator atau dielektrik. Ketika diberi tegangan, kapasitor menyimpan muatan pada kedua pelat dan energi pada medan listrik di antara pelat tersebut. Kapasitansi menunjukkan kemampuan menyimpan muatan; nilainya dipengaruhi oleh luas pelat, jarak antar-pelat, dan jenis bahan dielektrik. Kapasitor digunakan untuk penyaring catu daya, kopling sinyal, pelewat/penahan frekuensi tertentu, pewaktu, dan penyimpan energi sementara. Satuan kapasitansi adalah farad (F). Kapasitor polar, seperti elektrolit, wajib dipasang sesuai tanda + dan − serta batas tegangan kerja; kapasitor non-polar dapat dipasang bolak-balik pada rangkaian AC sesuai spesifikasinya.')
    add_capacitor_image(d)
    table(d,['Konversi satuan','Nilai setara'],[['1 F','1.000.000 µF'],['1 µF','1.000 nF = 1.000.000 pF'],['1 nF','1.000 pF'],['1 pF','0,001 nF']],[5,11])
    section(d,'3.2 Tabel Kode Kapasitor 3 Digit')
    table(d,['Kode','Pembacaan dalam pF','Hasil konversi'],[['101','10 × 10¹ pF','100 pF = 0,1 nF'],['102','10 × 10² pF','1.000 pF = 1 nF'],['103','10 × 10³ pF','10.000 pF = 10 nF = 0,01 µF'],['104','10 × 10⁴ pF','100.000 pF = 100 nF = 0,1 µF'],['222','22 × 10² pF','2.200 pF = 2,2 nF'],['472','47 × 10² pF','4.700 pF = 4,7 nF']],[3,7,6])
    para(d,'Pada kode tiga digit, dua digit pertama adalah angka penting dan digit ketiga menunjukkan jumlah nol dalam satuan pF. Huruf J, K, dan M menunjukkan toleransi ±5%, ±10%, dan ±20%.')
    section(d,'BAB IV — RANGKAIAN SERI DAN PARALEL')
    add_circuit_image(d)
    para(d,'Resistor: seri menggunakan penjumlahan langsung, Rs = R1 + R2 + ...; paralel menggunakan 1/Rp = 1/R1 + 1/R2 + ... atau untuk dua resistor Rp = (R1×R2)/(R1+R2).')
    para(d,'Kapasitor: paralel menggunakan penjumlahan langsung, Cp = C1 + C2 + ...; seri menggunakan 1/Cs = 1/C1 + 1/C2 + ... atau untuk dua kapasitor Cs = (C1×C2)/(C1+C2).')
    section(d,'Contoh Soal dan Pembahasan')
    numbered(d,['R1 = 1 kΩ dan R2 = 2 kΩ seri. Rs = 1 + 2 = 3 kΩ.','R1 = 2 kΩ dan R2 = 3 kΩ paralel. Rp = (2×3)/(2+3) = 1,2 kΩ.','C1 = 100 nF dan C2 = 220 nF paralel. Cp = 320 nF.','C1 = 100 nF dan C2 = 100 nF seri. Cs = (100×100)/(100+100) = 50 nF.'])
    section(d,'6. Rangkuman dan Keselamatan Kerja')
    bullets(d,['Baca kode dari arah yang benar dan tuliskan satuan hasil.','Periksa kembali faktor pengali dan konversi pF–nF–µF.','Jangan memasang kapasitor elektrolit terbalik; perhatikan batas tegangan kerja.','Gunakan komponen dan alat ukur dengan tertib, kering, dan sesuai petunjuk.'])
    section(d,'7. Latihan Mandiri')
    numbered(d,['Tentukan nilai resistor gelang kuning–ungu–cokelat–emas.','Ubah kode kapasitor 103 ke dalam nF dan µF.','Hitung Rs untuk 470 Ω dan 1,5 kΩ seri.','Hitung Cp untuk 47 nF dan 100 nF paralel.','Hitung Cs untuk dua kapasitor 220 nF yang disusun seri.'])
    para(d,'Kunci singkat: 1) 470 Ω ±5%; 2) 10 nF = 0,01 µF; 3) 1,97 kΩ; 4) 147 nF; 5) 110 nF.')
    return save(d,'BahanAjar_ResistorKapasitor_X-TEI_GegeDesembri.docx')

def make_modul():
    d=Document(); style_doc(d); common_identity(d,'MODUL AJAR')
    section(d,'I. INFORMASI UMUM')
    table(d,['Komponen','Isi'],[['Kompetensi awal','Peserta didik mengenal besaran listrik dasar dan operasi hitung sederhana.'],['Target peserta didik','Peserta didik reguler kelas X TEI; pembelajaran berkelompok heterogen.'],['Sarana/prasarana','Resistor, kapasitor, LKPD, papan tulis/LCD, kalkulator, multimeter bila tersedia.'],['Model pembelajaran','Problem Based Learning (PBL)'],['Kata kunci','Resistor, kapasitor, kode warna, kode angka, seri, paralel, toleransi.']],[4.5,11.5])
    section(d,'II. KOMPONEN INTI')
    para(d,'CP dan tujuan: peserta didik memahami komponen pasif, membaca kode nilai, dan menerapkan perhitungan resistor serta kapasitor dalam rangkaian sederhana sesuai TP DTE 9.1 dan DTE 9.3.')
    para(d,'Pemahaman bermakna: kemampuan membaca nilai komponen dan menghitung nilai pengganti membantu teknisi memilih, memeriksa, dan memperbaiki rangkaian elektronika secara aman dan teliti.')
    para(d,'Pertanyaan pemantik: Mengapa dua komponen yang bentuknya mirip dapat memiliki nilai berbeda? Mengapa rumus kapasitor seri berbeda dari resistor seri?')
    section(d,'III. KEGIATAN PEMBELAJARAN (240 MENIT)')
    table(d,['Tahap PBL','Aktivitas Guru dan Peserta Didik','Waktu'],[['Pendahuluan','Salam, presensi, apersepsi, tujuan, manfaat, dan kesepakatan kerja.','20'],['Orientasi masalah','Menghadirkan komponen R/C dan masalah identifikasi nilai serta pemilihan susunan rangkaian.','25'],['Mengorganisasi','Membentuk kelompok dan membagi LKPD/peran.','15'],['Membimbing penyelidikan','Membaca kode, mengonversi satuan, menghitung R/C seri-paralel, dan mendokumentasikan langkah.','95'],['Menyajikan hasil','Presentasi kelompok dan tanggapan.','35'],['Menganalisis/evaluasi','Validasi jawaban dan penguatan konsep.','20'],['Penutup','Tes individu, refleksi, kesimpulan, tindak lanjut, doa.','30']],[3.2,10.8,2])
    section(d,'IV. ASESMEN')
    bullets(d,['Diagnostik: pertanyaan awal tentang simbol dan satuan R/C.','Formatif: observasi proses, tanya jawab, pemeriksaan tabel kode dan langkah hitung pada LKPD.','Sumatif: tes individu 5 soal dan penilaian produk/presentasi kelompok.'])
    section(d,'V. DIFERENSIASI DAN TINDAK LANJUT')
    bullets(d,['Konten: kartu bantuan kode warna dan tabel konversi untuk peserta didik yang memerlukan dukungan.','Proses: tutor sebaya dan contoh bertahap; peserta didik cepat diberi soal rangkaian campuran.','Produk: hasil dapat disajikan sebagai tabel perhitungan atau penjelasan lisan terstruktur.','Remedial dan pengayaan mengikuti hasil asesmen; guru mencatat perkembangan pada jurnal pembelajaran.'])
    section(d,'VI. REFLEKSI')
    table(d,['Refleksi Peserta Didik','Refleksi Guru'],[['Konsep apa yang paling saya pahami?','Apakah alokasi 240 menit efektif?'],['Bagian mana yang masih membingungkan?','Miskonsepsi apa yang muncul?'],['Bagaimana kontribusi saya dalam kelompok?','Kelompok mana yang membutuhkan tindak lanjut?']],[8,8])
    section(d,'VII. LAMPIRAN')
    bullets(d,['LKPD Resistor dan Kapasitor.','Kunci jawaban dan rubrik penilaian.','Ringkasan rumus dan tabel kode.'])
    return save(d,'ModulAjar_ResistorKapasitor_X-TEI_GegeDesembri.docx')

def make_lkpd():
    d=Document(); style_doc(d); title(d,'LEMBAR KERJA PESERTA DIDIK (LKPD)','RESISTOR DAN KAPASITOR — PROBLEM BASED LEARNING')
    table(d,['Identitas','Isian'],[['Nama kelompok','........................................................'],['Anggota','1. ................ 2. ................ 3. ................ 4. ................'],['Kelas/Tanggal','X TEI / Jumat, 25 September 2026'],['Alokasi','6 JP × 40 menit']],[4.5,11.5])
    section(d,'A. TUJUAN')
    bullets(d,['Mengidentifikasi fungsi dan simbol resistor serta kapasitor.','Membaca kode nilai resistor dan kapasitor.','Menghitung rangkaian seri-paralel R dan C dengan langkah yang benar.','Menyajikan hasil secara teliti dan bertanggung jawab.'])
    section(d,'B. PETUNJUK KERJA')
    numbered(d,['Bekerjalah dalam kelompok sesuai peran yang disepakati.','Amati komponen dengan hati-hati dan tuliskan satuan setiap hasil.','Tunjukkan rumus, substitusi angka, dan kesimpulan pada setiap perhitungan.','Bandingkan hasil kelompok sebelum presentasi.'])
    section(d,'C. KEGIATAN 1 — IDENTIFIKASI DAN PEMBACAAN KODE')
    add_image(d, IMG, width_cm=12)
    para(d,'Lengkapi tabel berikut berdasarkan komponen contoh atau data yang diberikan guru.')
    table(d,['No','Komponen/kode','Jenis/fungsi','Nilai hasil baca','Toleransi/satuan'],[['1','Cokelat–Hitam–Merah–Emas','................','................','................'],['2','Kuning–Ungu–Cokelat–Emas','................','................','................'],['3','Kapasitor 104','................','................','................'],['4','Kapasitor 472','................','................','................'],['5','Komponen pilihan kelompok','................','................','................']],[1,4,4,4,3])
    section(d,'D. KEGIATAN 2 — PERHITUNGAN RANGKAIAN')
    add_circuit_image(d)
    numbered(d,['R1 = 470 Ω dan R2 = 1,5 kΩ disusun seri. Tentukan Rs.','R1 = 2 kΩ dan R2 = 3 kΩ disusun paralel. Tentukan Rp.','C1 = 100 nF dan C2 = 220 nF disusun paralel. Tentukan Cp.','C1 = 100 nF dan C2 = 100 nF disusun seri. Tentukan Cs.','Tantangan: tiga kapasitor 10 nF, 20 nF, dan 30 nF disusun paralel. Tentukan Ct dan jelaskan manfaat susunan paralel.'])
    for i in range(5):
        para(d,'Ruang jawaban '+str(i+1)+':\n\n........................................................................................................................\n........................................................................................................................\n........................................................................................................................',size=11)
    section(d,'E. KESIMPULAN DAN REFLEKSI')
    para(d,'Tuliskan dua kesimpulan konsep dan satu hal yang masih perlu dipelajari.\n1. ........................................................................................................................\n2. ........................................................................................................................\nHal yang perlu dipelajari: ...............................................................................................',size=11)
    section(d,'F. RUBRIK PENILAIAN')
    table(d,['Aspek','4','3','2','1'],[['Pembacaan kode','Semua tepat','Ada 1 kesalahan','Ada 2 kesalahan','Belum mampu'],['Perhitungan','Rumus dan hasil tepat','Ada kesalahan kecil','Langkah belum lengkap','Tidak menunjukkan langkah'],['Kerja sama','Aktif dan berbagi peran','Cukup aktif','Kurang merata','Tidak bekerja sama'],['Presentasi','Jelas, sistematis, menjawab','Cukup jelas','Kurang sistematis','Tidak mampu menjelaskan']],[3.5,3.1,3.1,3.1,3.1])
    return save(d,'LKPD_ResistorKapasitor_X-TEI_GegeDesembri.docx')

def make_kunci():
    d=Document(); style_doc(d); title(d,'KUNCI JAWABAN','LKPD RESISTOR DAN KAPASITOR — KELAS X TEI')
    para(d,'Guru: '+TEACHER+' | SMKN 1 Terusan Nunyai | Alokasi: 6 JP × 40 menit')
    section(d,'A. KUNCI KEGIATAN 1 — IDENTIFIKASI DAN PEMBACAAN KODE')
    table(d,['No','Komponen/kode','Jawaban'],[['1','Cokelat–Hitam–Merah–Emas','1 kΩ ±5%'],['2','Kuning–Ungu–Cokelat–Emas','470 Ω ±5%'],['3','Kapasitor 104','100 nF = 0,1 µF'],['4','Kapasitor 472','4,7 nF'],['5','Komponen pilihan kelompok','Sesuai komponen yang dipilih dan diverifikasi guru.']],[1,6,9])
    section(d,'B. KUNCI KEGIATAN 2 — PERHITUNGAN RANGKAIAN')
    table(d,['No','Soal','Jawaban'],[['1','470 Ω dan 1,5 kΩ seri','1,97 kΩ'],['2','2 kΩ dan 3 kΩ paralel','1,2 kΩ'],['3','100 nF dan 220 nF paralel','320 nF'],['4','100 nF dan 100 nF seri','50 nF'],['5','10 nF, 20 nF, dan 30 nF paralel','60 nF']],[1,8,7])
    section(d,'C. PEDOMAN PENSKORAN')
    table(d,['Aspek','Skor maksimum'],[['Pembacaan kode','20'],['Perhitungan R dan C','40'],['Kerja sama','20'],['Presentasi','20']],[12,4])
    para(d,'Catatan: nilai akhir disesuaikan dengan rubrik proses dan presentasi pada LKPD.',size=11)
    return save(d,'KunciJawaban_LKPD_ResistorKapasitor_X-TEI_GegeDesembri.docx')

if __name__=='__main__':
    results=[make_rpp(),make_bahan(),make_modul(),make_lkpd(),make_kunci()]
    for a,b in results: print(a); print(b)
